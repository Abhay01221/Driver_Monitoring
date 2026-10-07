from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

INPUT = Path(r"C:\Users\abhay\Downloads\CSE411_Task4_Driver_Drowsiness_Detection.docx")
OUTPUT = Path(r"C:\Users\abhay\Driver_Monitoring\output\docx\CSE411_Task4_Driver_Drowsiness_Detection_completed.docx")
ROOT = Path(r"C:\Users\abhay\Driver_Monitoring")

def shade(paragraph, fill="F3F6F9"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "4"); el.set(qn("w:color"), "D7E0E8"); borders.append(el)
    ppr.append(borders)

def add_code(doc, title, rel_path, start, end):
    p = doc.add_paragraph(); p.add_run(title).bold = True
    src = doc.add_paragraph(); src.paragraph_format.space_after = Pt(3)
    run = src.add_run(f"Source: {rel_path}"); run.italic = True; run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(92, 107, 120)
    lines = (ROOT / rel_path).read_text(encoding="utf-8").splitlines()
    code = doc.add_paragraph(); code.paragraph_format.left_indent = Inches(.08)
    code.paragraph_format.right_indent = Inches(.08); code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(8); code.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    shade(code)
    for i, line in enumerate(lines[start-1:end], start):
        r = code.add_run(f"{i:>3}  {line}\n"); r.font.name = "Courier New"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New"); r.font.size = Pt(7.3)
        r.font.color.rgb = RGBColor(31, 41, 55)

doc = Document(str(INPUT))
doc.add_page_break()
h = doc.add_paragraph(); run = h.add_run("6. Code Implementation from the Project"); run.bold = True; run.font.size = Pt(16); run.font.color.rgb = RGBColor(20, 40, 61)
doc.add_paragraph("The following excerpts are taken from the implemented repository and correspond to the modules described in the reference report. They show the working dataset-to-inference path and the temporal alarm behavior.")
add_code(doc, "Code Snippet 1 - Webcam Frame Capture", r"frontend/components/Camera.jsx", 64, 86)
add_code(doc, "Code Snippet 2 - Frontend API Request", r"frontend/lib/api.js", 28, 57)
add_code(doc, "Code Snippet 3 - Image Validation and Prediction Route", r"backend/app/routes/predict.py", 24, 82)
add_code(doc, "Code Snippet 4 - Model Loading and Class Safeguard", r"backend/app/models/model.py", 25, 55)
add_code(doc, "Code Snippet 5 - Five Second Drowsiness Alarm", r"frontend/app/page.jsx", 76, 119)
doc.add_paragraph("Implementation note: place the optional alarm audio at frontend/public/alarm.mp3. Live monitoring evaluates roughly one frame per second and starts the alarm after five seconds of continuous Drowsy predictions.")
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUTPUT))
print(OUTPUT.resolve())
